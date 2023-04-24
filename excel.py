
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import PatternFill
from openpyxl.styles import Font, Color, Alignment, Border, Side

from classes import Stats
from mapping import FRAGMENTATION_INDEX, SLEEP_LATENCY, ACTUAL_WAKE_TIME, SLEEP_EFFICIENCY, LIGHTS_OUT, GOT_UP, TIME_IN_BED, ASSUMED_SLEEP, ACTUAL_SLEEP_TIME
from mapping import USERID, ACTUAL_SLEEP_RATE, ACTUAL_WAKE_RATE, FELL_ASLEEP, WOKE_UP


stats = []

workbook2 = load_workbook(filename = "donnees_actimetres.xlsx", read_only=True)
workbook2.sheetnames

#parcours des feuilles
for sheet in workbook2 : 

    #les titres des feuilles sont au format USERID_TEMP (c'est moi qui l'ai décidé nah)
    #  --> n'ayant pas de vraibale température sur MotionWare on récupère la température de la nuit à partir du nom de la feuille
    temperaure = int(sheet.title[-2:])

    #remplissage de la classe
    stat = Stats(id = sheet[USERID].value, temp = temperaure, assumed_sleep=sheet[ASSUMED_SLEEP].value, actual_sleep_time= sheet[ACTUAL_SLEEP_TIME].value, 
                 actual_sleep_rate = sheet[ACTUAL_SLEEP_RATE].value, actual_wake_time=sheet[ACTUAL_WAKE_TIME].value, actual_wake_rate = sheet[ACTUAL_WAKE_RATE].value, 
                 time_in_bed=sheet[TIME_IN_BED].value, sleep_efficiency=sheet[SLEEP_EFFICIENCY].value,lights_out=sheet[LIGHTS_OUT].value, fell_asleep=sheet[FELL_ASLEEP].value, 
                 sleep_latency=sheet[SLEEP_LATENCY].value ,woke_up=sheet[WOKE_UP].value, got_up=sheet[GOT_UP].value, fragmentation_index=sheet[FRAGMENTATION_INDEX].value)
    
    stats.append(stat)
    #print(stats)

workbook = Workbook()
sheet = workbook.active

#nom des en-têtes
sheet.append(["UserID", "temperature", "time_in_bed", "assumed_sleep", "actual_sleep_time", "actual_sleep (%)", "actual_wake_time", "actual_wake (%)", "sleep_efficiency (%)", "lights_out", 
              "fell_asleep", "sleep_latency", "woke_up", "got_up", "fragmentation_index"])

#remplissage du tableau
for stat in stats : 
    data = [stat.id, stat.temp, stat.time_in_bed, stat.assumed_sleep, stat.actual_sleep_time, stat.actual_sleep_rate, stat.actual_wake_time, stat.actual_wake_rate,stat.sleep_efficiency,
            stat.lights_out, stat.fell_asleep, stat.sleep_latency, stat.woke_up, stat.got_up, stat.fragmentation_index]
    sheet.append(data)

#Ajustement de la largeur de la colonne en fonction de son contenu
dims = {}
for row in sheet.rows:
    for cell in row:
        if cell.value:
             dims[cell.column_letter] = max((dims.get(cell.column_letter, 0), len(str(cell.value))))  
for col, value in dims.items():
    sheet.column_dimensions[col].width = value

#remplissage de ligne alternatif
gris = "e2e2e2"
for rows in sheet.iter_rows(min_row=2, max_row=69, min_col=1, max_col=15):
    for cell in rows:
        if cell.row % 2:
            cell.fill = PatternFill(start_color=gris, end_color=gris,fill_type = "solid")

#code couleur pour la température
color_scale_rule = ColorScaleRule(end_type="num",
                                  end_value=32,
                                  end_color="e67d59",  # Rouge
                                  mid_type="num",
                                  mid_value=24,
                                  mid_color="9edd87",  # vert
                                  start_type="num",
                                  start_value=16,
                                  start_color="c3d2ff")  # bleu
sheet.conditional_formatting.add("B2:B69", color_scale_rule)


#filtres
sheet.auto_filter.ref = sheet.dimensions

#epingle la première ligne
sheet.freeze_panes = "B2"

workbook.save(filename= "test.xlsx")